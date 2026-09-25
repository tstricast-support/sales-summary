import io, os
from contextlib import asynccontextmanager
from datetime import date
from fastapi import FastAPI, Depends, HTTPException, Header, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy.orm import Session
import crud, models, schemas
from database import engine, Base, get_db, SessionLocal

ADMIN_KEY = os.getenv("ADMIN_KEY", "")


@asynccontextmanager
async def lifespan(app):
    Base.metadata.create_all(bind=engine)
    with SessionLocal() as db:
        crud.seed_departments(db)
    yield


app = FastAPI(title="Summary API", lifespan=lifespan)
app.add_middleware(CORSMiddleware, allow_origins=[o.strip() for o in os.getenv("CORS_ORIGINS", "*").split(",")],
                   allow_methods=["*"], allow_headers=["*"])


def admin_only(x_admin_key: str = Header(default="")):
    if ADMIN_KEY and x_admin_key != ADMIN_KEY:
        raise HTTPException(401, "Admin key required")


def check_range(start: date, end: date):
    if start > end:
        raise HTTPException(422, "Start date must be before end date")


@app.get("/api/health")
def health():
    return {"status": "ok"}


@app.get("/api/departments", response_model=list[schemas.DepartmentOut])
def departments(db: Session = Depends(get_db)):
    return db.query(models.Department).order_by(models.Department.id).all()


@app.get("/api/records", response_model=list[schemas.RecordOut])
def records(start: date, end: date, department: str | None = None, db: Session = Depends(get_db)):
    check_range(start, end)
    return crud.list_records(db, start, end, department)


@app.post("/api/records", response_model=schemas.RecordOut)
def save_record(data: schemas.RecordIn, db: Session = Depends(get_db)):
    try:
        return crud.upsert_record(db, data)
    except LookupError as e:
        raise HTTPException(404, str(e))

@app.get("/api/damages", response_model=list[schemas.DamageOut])
def damages(db: Session = Depends(get_db)):
    return crud.list_damages(db)


@app.post("/api/damages", response_model=schemas.DamageOut)
def add_damage(data: schemas.DamageIn, db: Session = Depends(get_db)):
    return crud.add_damage(db, data)

@app.post("/api/records/bulk")
def save_bulk(data: schemas.BulkIn, db: Session = Depends(get_db)):
    try:
        return {"days": crud.bulk_upsert(db, data)}
    except LookupError as e:
        raise HTTPException(404, str(e))
    except ValueError as e:
        raise HTTPException(422, str(e))

@app.get("/api/summary", dependencies=[Depends(admin_only)])
def summary(start: date, end: date, db: Session = Depends(get_db)):
    check_range(start, end)
    return crud.summary(db, start, end)


@app.get("/api/audit", dependencies=[Depends(admin_only)])
def audit(start: date, end: date, department: str | None = None, db: Session = Depends(get_db)):
    check_range(start, end)
    return crud.audit(db, start, end, department)


@app.get("/api/export", dependencies=[Depends(admin_only)])
def export(start: date, end: date, department: str | None = None, db: Session = Depends(get_db)):
    check_range(start, end)
    rows = crud.list_records(db, start, end, department)
    wb = Workbook(); ws = wb.active; ws.title = "Sales & Collections"
    ws.append(["Date", "Department", "Sales", "Collection", "Submitted by", "Last updated"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([r.record_date, r.department.name, float(r.sales_amount), float(r.collection_amount),
                   r.submitted_by, r.updated_at.strftime("%Y-%m-%d %H:%M") if r.updated_at else ""])
    n = len(rows) + 1
    ws.append(["TOTAL", "", f"=SUM(C2:C{n})", f"=SUM(D2:D{n})"])
    for c in ws[n + 1]:
        c.font = Font(bold=True)
    for col, w in zip("ABCDEF", (12, 20, 16, 16, 20, 18)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=4):
        for c in row:
            c.number_format = "#,##0.00"
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    name = f"summary_{department or 'all'}_{start}_{end}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{name}"'})
